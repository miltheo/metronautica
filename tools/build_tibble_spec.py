"""Build the Word copy from the maintained Metronautica Tibble Markdown.

Requires Python and python-docx. The companion PowerShell script also exports
PDF through Microsoft Word. No package calculations or R dependencies are used.
"""
from pathlib import Path
import argparse
import re
import textwrap

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
VERSION = "0.1.0"
SOURCE = ROOT / "data-contract" / VERSION / "SPECIFICATION.md"


def font_properties(properties, family):
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for key in list(fonts.attrib):
        del fonts.attrib[key]
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), family)


def style(doc, name, size, *, family="Calibri", bold=False, before=0, after=6):
    value = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    value.font.name = family
    value.font.size = Pt(size)
    value.font.bold = bold
    value.font.italic = False
    value.font.underline = False
    value.font.all_caps = False
    value.font.small_caps = False
    value.font.color.rgb = RGBColor.from_string("000000")
    font_properties(value.element.get_or_add_rPr(), family)
    value.paragraph_format.space_before = Pt(before)
    value.paragraph_format.space_after = Pt(after)
    value.paragraph_format.line_spacing = 1.10
    value.paragraph_format.widow_control = True
    return value


def field(paragraph, instruction):
    run = paragraph.add_run()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = " " + instruction + " "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (start, code, end):
        run._r.append(element)


def inline(paragraph, value, *, code=False, size=None):
    run = paragraph.add_run(value)
    if code:
        run.font.name = "Consolas"
        run.font.size = Pt(size or 10)
        font_properties(run._r.get_or_add_rPr(), "Consolas")
        # Disable proofing of literal identifiers and code.
        run._r.get_or_add_rPr().append(OxmlElement("w:noProof"))
    return run


TOKEN = re.compile(r"\b(?:[A-Za-z][A-Za-z0-9]*_[A-Za-z0-9_]+(?:\(\))?|POSIXct|tbl_df)(?:\$[A-Za-z_][A-Za-z0-9_]*)*|\bmeta(?:\$[A-Za-z_][A-Za-z0-9_]*)+")


def prose(paragraph, value):
    end = 0
    for match in TOKEN.finditer(value):
        inline(paragraph, value[end:match.start()])
        inline(paragraph, match.group(), code=True)
        end = match.end()
    inline(paragraph, value[end:])


def table_cells(line):
    return [cell.strip().replace(r"\|", "|").replace("<br>", "\n")
        for cell in re.split(r"(?<!\\)\|", line.strip())[1:-1]]


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            yield "code", "\n".join(code)
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:|\-]+\|$", lines[i]):
                    rows.append(table_cells(lines[i]))
                i += 1
            yield "table", rows
        elif line.startswith("#"):
            heading = re.match(r"^(#+) (.*)$", line)
            yield "heading", (len(heading[1]), heading[2])
            i += 1
        else:
            paragraph = []
            while i < len(lines) and lines[i].strip():
                paragraph.append(lines[i])
                i += 1
            yield "paragraph", " ".join(paragraph)


def add_table(doc, rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    widths = (2.05, 1.00, 3.85) if len(rows[0]) == 3 else (2.1, 1.6, 1.6, 1.6)
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + key), value)
        borders.append(border)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, size in (("top", 65), ("bottom", 65), ("left", 95), ("right", 95)):
        margin = OxmlElement("w:" + side)
        margin.set(qn("w:w"), str(size))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E9EDF1" if index == 0 else "FFFFFF" if index % 2 else "F7F8FA")
            cell._tc.get_or_add_tcPr().append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["Table Header" if index == 0 else "Table Text"]
            if column == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if index and column == 0:
                value = "\n".join(textwrap.wrap(value, width=27,
                    break_long_words=False, break_on_hyphens=False))
                inline(paragraph, value, code=True, size=9.25)
            else:
                inline(paragraph, value)
            # The short core-column table stays together.
            if len(rows) == 5 and index < len(rows) - 1:
                paragraph.paragraph_format.keep_with_next = True
    gap = doc.add_paragraph(style="Table Gap")
    gap.paragraph_format.keep_with_next = False
    return table


def build(output):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.78)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.header_distance = sec.footer_distance = Inches(0.35)
    normal = style(doc, "Normal", 11, after=7)
    normal.paragraph_format.keep_together = True
    style(doc, "Title", 25, bold=True, after=8)
    style(doc, "Subtitle", 11.5, after=8)
    style(doc, "Version", 9.5, after=18)
    for name, size, before, after in (("Heading 1", 15, 16, 7), ("Heading 2", 11.5, 12, 5), ("Heading 3", 11, 9, 4)):
        heading = style(doc, name, size, bold=True, before=before, after=after)
        heading.paragraph_format.keep_with_next = True
    style(doc, "Table Header", 9.5, bold=True, after=0)
    style(doc, "Table Text", 9.5, after=0).paragraph_format.line_spacing = 1.06
    style(doc, "Table Gap", 2, after=4)
    code_style = style(doc, "Code Block", 10, family="Consolas", before=7, after=11)
    code_style.paragraph_format.left_indent = Inches(0.12)
    code_style.paragraph_format.right_indent = Inches(0.12)
    code_style.paragraph_format.line_spacing = 1.18
    code_style.paragraph_format.keep_together = True
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F3F5F7")
    code_style.element.get_or_add_pPr().append(shade)
    style(doc, "Running Text", 8.5, after=0)

    # Remove decorations and character spacing inherited from Word templates.
    for value in doc.styles:
        for element in list(value.element.iter(qn("w:pBdr"))):
            element.getparent().remove(element)
        for element in list(value.element.iter(qn("w:spacing"))):
            if element.getparent().tag == qn("w:rPr"):
                element.getparent().remove(element)

    header = sec.header.paragraphs[0]
    header.style = doc.styles["Running Text"]
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT)
    header.add_run("METRONAUTICA TIBBLE\tSPECIFICATION " + VERSION)
    footer = sec.footer.paragraphs[0]
    footer.style = doc.styles["Running Text"]
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run("Metronautica · " + VERSION + "\t")
    field(footer, "PAGE")
    footer.add_run(" / ")
    field(footer, "NUMPAGES")

    paragraph_number = 0
    for kind, value in blocks(SOURCE.read_text(encoding="utf-8")):
        if kind == "heading":
            level, title = value
            paragraph = doc.add_paragraph(title, "Title" if level == 1 else f"Heading {level - 1}")
            if level == 1:
                paragraph.paragraph_format.keep_with_next = True
            elif level == 2 and title.startswith(("2 ", "3 ", "4 ", "5 ")):
                paragraph.paragraph_format.page_break_before = True
        elif kind == "paragraph":
            if value.startswith("[Specification index]"):
                # Keep navigation in repository guides.
                continue
            paragraph_number += 1
            if paragraph_number == 1:
                paragraph = doc.add_paragraph(value, "Subtitle")
                paragraph.paragraph_format.keep_with_next = True
            elif paragraph_number == 2:
                paragraph = doc.add_paragraph(value, "Version")
                paragraph.paragraph_format.keep_with_next = True
            elif value == "Minimal construction without metadata":
                doc.add_paragraph(value, "Heading 3")
            else:
                paragraph = doc.add_paragraph()
                prose(paragraph, value)
        elif kind == "code":
            paragraph = doc.add_paragraph(style="Code Block")
            inline(paragraph, value, code=True, size=10)
        elif kind == "table":
            add_table(doc, value)

    doc.core_properties.title = "Metronautica Tibble"
    doc.core_properties.subject = "Specification 0.1.0 with optional Metronautica pipeline fields"
    doc.core_properties.author = "Metronautica"
    doc.core_properties.keywords = "acceleration, data contract, Metronautica Tibble"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=SOURCE.with_name("Metronautica_Tibble_0.1.0.docx"))
    build(parser.parse_args().output)
